"""Curated dataset specs and the parsing of their resources (CSV or XLSX) into compact tables."""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path
from typing import Any

import openpyxl
import yaml

SPEC_PATH = Path(__file__).with_name("datasets.yaml")
MAX_ROWS = 5000


class TableError(Exception):
    pass


@dataclass(frozen=True)
class DatasetSpec:
    key: str
    portal: str
    title: str
    publisher: str
    url: str
    resource: dict[str, Any]
    table: dict[str, Any]
    services: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class Catalogue:
    datasets: list[DatasetSpec]
    related: dict[str, list[str]]

    def by_key(self, key: str) -> DatasetSpec | None:
        return next((d for d in self.datasets if d.key == key), None)

    def for_service(self, service_id: str) -> list[DatasetSpec]:
        return [d for d in self.datasets if service_id in d.services]


@lru_cache
def catalogue() -> Catalogue:
    raw = yaml.safe_load(SPEC_PATH.read_text(encoding="utf-8"))
    return Catalogue(
        datasets=[DatasetSpec(**item) for item in raw["datasets"]],
        related={k: list(v) for k, v in (raw.get("related") or {}).items()},
    )


def norm(text: Any) -> str:
    """Header names compare without case, line breaks or repeated spaces."""
    return " ".join(str(text or "").split()).lower()


def pick_resource(spec: DatasetSpec, resources: list[Any]) -> Any:
    """Choose the resource described by spec.resource from a dataset page's resource list."""
    want = spec.resource
    for resource in resources:
        name = resource.name.lower()
        if want.get("format") and resource.format != want["format"]:
            continue
        if want.get("name_contains") and want["name_contains"].lower() not in name:
            continue
        if any(word.lower() in name for word in want.get("name_excludes", [])):
            continue
        return resource
    raise TableError(f"{spec.key}: no resource matches {want} on the dataset page")


def _cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        # Keep list cells readable: one item per line in the source becomes "a; b; c".
        lines = [" ".join(part.split()) for part in value.splitlines() if part.strip()]
        text = "; ".join(lines)
        if re.fullmatch(r"-?\d+(\.\d+)?", text):
            number = float(text)
            return int(number) if number.is_integer() else number
        return text
    return value


def _rows(content: bytes, spec: DatasetSpec) -> list[list[Any]]:
    fmt = spec.table["format"]
    if fmt == "csv":
        text = content.decode("utf-8-sig", errors="replace")
        return [list(r) for r in csv.reader(io.StringIO(text))]
    if fmt == "xlsx":
        try:
            book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 - any parse failure means an unusable file
            raise TableError(
                f"{spec.key}: not a readable spreadsheet ({type(exc).__name__})"
            ) from exc
        sheet_name = spec.table.get("sheet")
        sheet = (
            book[sheet_name] if sheet_name and sheet_name in book.sheetnames else book.worksheets[0]
        )
        return [list(r) for r in sheet.iter_rows(values_only=True)]
    raise TableError(f"{spec.key}: unsupported table format {fmt!r}")


def parse_table(content: bytes, spec: DatasetSpec) -> tuple[list[str], list[list[Any]]]:
    """Header row found by `header_contains`; data rows until the first empty row; only `keep`."""
    rows = _rows(content, spec)
    marker = norm(spec.table["header_contains"])
    try:
        header_index = next(i for i, row in enumerate(rows) if any(marker in norm(c) for c in row))
    except StopIteration as exc:
        raise TableError(f"{spec.key}: header {spec.table['header_contains']!r} not found") from exc
    header = [" ".join(str(c or "").split()) for c in rows[header_index]]
    keep = spec.table.get("keep")
    if keep:
        indexes = []
        for name in keep:
            try:
                indexes.append(next(i for i, h in enumerate(header) if norm(h) == norm(name)))
            except StopIteration as exc:
                raise TableError(f"{spec.key}: column {name!r} not found") from exc
    else:
        indexes = [i for i, h in enumerate(header) if h]
    columns = [header[i] for i in indexes]
    data: list[list[Any]] = []
    for row in rows[header_index + 1 :]:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            break
        data.append([_cell(row[i]) if i < len(row) else None for i in indexes])
        if len(data) >= MAX_ROWS:
            break
    if not data:
        raise TableError(f"{spec.key}: the table has no rows")
    return columns, data
